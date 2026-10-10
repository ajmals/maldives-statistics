import os
import re
import json
import csv
import glob
import time
import openpyxl
from openpyxl.utils import get_column_letter

THAANA_REGEX = re.compile(r'[\u0780-\u07BF]')

LEGACY_FONTS = {'faseyha', 'faruma', 'randhoo', 'kuredhi', 'waheed', 'voguepsmt', 'a_faseyha', 'a_faruma', 'a_randhoo'}

LEGACY_WORDS = {
    'clwvwt', 'ulwvwt', 'cnunemihob', 'cscnijed', 'cnehirif', 'cnehcnwa', 
    'wlcmuj', 'cnwt', 'elwm', 'cnwg', 'udwminwh', 'udcawk', 'udcaedwk',
    'idwbwa', 'cswdcnurwhwa', 'wrufurumua', 'cswkwhwvudurwhwa', 'urwvcnim',
    'ctcneswp', 'ckirwt', 'urwhwa'
}

def is_dhivehi_font(font_name):
    if not font_name:
        return False
    fn = font_name.lower()
    return any(lf in fn for lf in LEGACY_FONTS)

def is_pure_dhivehi_cell(val, font_name=''):
    if val is None:
        return False
    val_str = str(val).strip()
    if not val_str:
        return False
    if isinstance(val, (int, float)) or re.match(r'^-?[\d,.]+$', val_str):
        return False
    if THAANA_REGEX.search(val_str):
        if not re.search(r'[A-Za-z]{3,}', val_str):
            return True
    if is_dhivehi_font(font_name):
        return True
    words = re.findall(r'[a-zA-Z]+', val_str.lower())
    if words and all(w in LEGACY_WORDS for w in words):
        return True
    return False

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        val_rounded = round(val, 8)
        val_str = f"{val_rounded:f}".rstrip('0').rstrip('.')
        return val_str
    if isinstance(val, int):
        return str(val)
    val_str = str(val).strip()
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

def clean_english_cell_text(val):
    if not val:
        return ""
    val_str = clean_val(val)
    val_str = THAANA_REGEX.sub('', val_str).strip()
    pattern = r'\b(' + '|'.join(LEGACY_WORDS) + r')\b'
    val_str = re.sub(pattern, '', val_str, flags=re.IGNORECASE).strip()
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

def is_multi_year_subtitle(txt):
    if not txt:
        return False
    return bool(re.search(r'\d{4}\s*,\s*\d{4}', txt) or re.search(r'\d{4}\s+AND\s+\d{4}', txt, re.IGNORECASE))

def process_single_table(xlsx_path, output_dir):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    base_name = os.path.splitext(os.path.basename(xlsx_path))[0].strip()
    base_name_clean = base_name.replace(' ', '')
    
    # 1. Target sheet selection
    target_sheet = None
    for name in wb.sheetnames:
        sheet = wb[name]
        if sheet.sheet_state != 'hidden':
            if name.strip().replace(' ', '') == base_name_clean:
                target_sheet = sheet
                break
    if not target_sheet:
        for name in wb.sheetnames:
            sheet = wb[name]
            if sheet.sheet_state != 'hidden':
                target_sheet = sheet
                break
                
    sheet = target_sheet
    
    # Bounding
    min_col, min_row = 1, 1
    max_col, max_row = sheet.max_column, sheet.max_row
    
    if sheet.print_area:
        pa = sheet.print_area
        if isinstance(pa, list) and pa:
            pa = pa[0]
        m = re.search(r'\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)', str(pa))
        if m:
            c1, r1, c2, r2 = m.groups()
            min_col = openpyxl.utils.column_index_from_string(c1)
            min_row = int(r1)
            max_col = openpyxl.utils.column_index_from_string(c2)
            max_row = int(r2)

    # 2. Build 2D grid
    grid = [[None for _ in range(max_col - min_col + 1)] for _ in range(max_row - min_row + 1)]
    font_grid = [["" for _ in range(max_col - min_col + 1)] for _ in range(max_row - min_row + 1)]
    
    for r_idx, r in enumerate(range(min_row, max_row + 1)):
        for c_idx, c in enumerate(range(min_col, max_col + 1)):
            cell = sheet.cell(r, c)
            grid[r_idx][c_idx] = cell.value
            font_grid[r_idx][c_idx] = cell.font.name if cell.font and cell.font.name else ""

    # Propagate merged cell values
    for rng in sheet.merged_cells.ranges:
        r_start = max(min_row, rng.min_row) - min_row
        r_end = min(max_row, rng.max_row) - min_row
        c_start = max(min_col, rng.min_col) - min_col
        c_end = min(max_col, rng.max_col) - min_col
        
        top_val = sheet.cell(rng.min_row, rng.min_col).value
        top_font = sheet.cell(rng.min_row, rng.min_col).font
        fname = top_font.name if top_font and top_font.name else ""
        
        for r_i in range(r_start, r_end + 1):
            for c_i in range(c_start, c_end + 1):
                if 0 <= r_i < len(grid) and 0 <= c_i < len(grid[0]):
                    grid[r_i][c_i] = top_val
                    font_grid[r_i][c_i] = fname

    # 3. Title, Metadata, Notes, Source
    title_en = ""
    title_dv = ""
    unit_info = ""
    notes = []
    source = ""
    data_end_row = len(grid)
    
    for r_idx, row in enumerate(grid):
        non_empty = [(c, clean_val(v), font_grid[r_idx][c]) for c, v in enumerate(row) if v is not None and clean_val(v) != ""]
        if not non_empty:
            continue
        first_c, first_v, first_f = non_empty[0]
        
        if re.search(r'Table\s+\d+(\.\d+)?[a-z]?\s*:', first_v, re.IGNORECASE):
            title_en = first_v
            continue
        if re.search(r'ތާވަލ[ުް]?\s+\d+(\.\d+)?[a-z]?\s*:', first_v) or first_v.endswith('clwvWt') or 'clwvWt :' in first_v or 'ulwvWt :' in first_v:
            title_dv = first_v
            continue
        if re.match(r'^\([A-Za-z0-9\s,%_/\-]+\)$', first_v) and not title_en:
            unit_info = first_v
            continue
            
        lower_v = first_v.lower()
        if any(kw in lower_v for kw in ['source:', 'source :', 'source ;', 'ﻤﻌﻠﻮﻣﺎﺕ', 'މައުލޫމާތު']):
            source = first_v
            if data_end_row == len(grid):
                data_end_row = r_idx
            continue
        if any(lower_v.startswith(kw) for kw in ['note:', 'note :', '* note', '1_/', '2_/', '3_/', '4_/', '5_/']):
            notes.append(first_v)
            if data_end_row == len(grid):
                data_end_row = r_idx
            continue
        if data_end_row < len(grid):
            notes.append(first_v)

    if not title_en:
        for r_idx in range(min(5, len(grid))):
            for c_idx, v in enumerate(grid[r_idx]):
                if v and isinstance(v, str) and 'Table' in v:
                    title_en = clean_val(v)
                    break
            if title_en:
                break

    # 4. Determine Header start & Data start
    header_start_row = -1
    data_start_row = -1
    
    for r_idx in range(min(15, data_end_row)):
        row = grid[r_idx]
        cells = [clean_val(v) for v in row if v is not None and clean_val(v) != ""]
        if not cells:
            continue
        line_str = " ".join(cells)
        if title_en and title_en in line_str:
            continue
        if title_dv and title_dv in line_str:
            continue
        if len(cells) == 1 and line_str.startswith('(') and line_str.endswith(')'):
            continue
        if is_multi_year_subtitle(line_str) and len(cells) == 1:
            continue
            
        nums = [v for v in row if v is not None and (isinstance(v, (int, float)) or re.match(r'^-?[\d,.]+$', clean_val(v)))]
        all_years = False
        if nums:
            all_years = all((isinstance(v, int) and 1900 <= v <= 2050) or (isinstance(v, str) and re.match(r'^(19|20)\d\d$', v.strip())) for v in nums)
            
        if nums and not all_years:
            data_start_row = r_idx
            break
        else:
            if header_start_row == -1 and len(cells) >= 1:
                header_start_row = r_idx

    if data_start_row == -1:
        data_start_row = min(header_start_row + 1, data_end_row)
    if header_start_row == -1:
        header_start_row = 0

    header_rows = grid[header_start_row:data_start_row]
    header_fonts = font_grid[header_start_row:data_start_row]
    raw_data_rows = grid[data_start_row:data_end_row]
    raw_data_fonts = font_grid[data_start_row:data_end_row]

    # Filter out pure Dhivehi data rows
    data_rows = []
    data_font_rows = []
    for r_idx, r in enumerate(raw_data_rows):
        has_num = any(isinstance(v, (int, float)) or (isinstance(v, str) and re.match(r'^-?[\d,.]+$', v.strip())) for v in r if v is not None)
        non_empty = [v for v in r if v is not None and clean_val(v) != ""]
        if not non_empty:
            continue
        if not has_num:
            all_dv = all(is_pure_dhivehi_cell(v, raw_data_fonts[r_idx][c]) for c, v in enumerate(r) if v is not None and clean_val(v) != "")
            if all_dv:
                continue
        data_rows.append(r)
        data_font_rows.append(raw_data_fonts[r_idx])

    num_cols = len(grid[0]) if grid else 0

    # Determine candidate columns
    keep_cols = []
    for c in range(num_cols):
        col_data = [r[c] for r in data_rows if r[c] is not None and clean_val(r[c]) != ""]
        col_header = [clean_val(r[c]) for r in header_rows if r[c] is not None and clean_val(r[c]) != ""]
        if not col_data and not col_header:
            continue
            
        all_col_cells = [(r[c], data_font_rows[i][c]) for i, r in enumerate(data_rows) if r[c] is not None and clean_val(r[c]) != ""]
        all_col_cells += [(r[c], header_fonts[i][c]) for i, r in enumerate(header_rows) if r[c] is not None and clean_val(r[c]) != ""]
        
        is_col_dhivehi = True
        for val, fname in all_col_cells:
            v_str = clean_val(val)
            if isinstance(val, (int, float)) or re.match(r'^-?[\d,.]+$', v_str):
                is_col_dhivehi = False
                break
            if not is_pure_dhivehi_cell(val, fname):
                is_col_dhivehi = False
                break
                
        if not is_col_dhivehi:
            keep_cols.append(c)

    # Drop duplicate consecutive columns
    filtered_keep_cols = []
    for i, c in enumerate(keep_cols):
        if i > 0:
            prev_c = keep_cols[i - 1]
            prev_vals = [clean_val(r[prev_c]) for r in data_rows]
            curr_vals = [clean_val(r[c]) for r in data_rows]
            if prev_vals and curr_vals and prev_vals == curr_vals:
                continue
        filtered_keep_cols.append(c)
    keep_cols = filtered_keep_cols

    # Forward-fill header labels across columns in keep_cols (excluding col 0)
    expanded_header_rows = []
    for h_idx, h_row in enumerate(header_rows):
        exp_row = [None] * num_cols
        last_val = None
        for c in range(num_cols):
            if c > 0 and (c - 1) not in keep_cols:
                last_val = None
                
            val = h_row[c]
            fname = header_fonts[h_idx][c]
            if val is not None and not is_pure_dhivehi_cell(val, fname):
                txt = clean_english_cell_text(val)
                if title_en and txt.lower() == title_en.lower():
                    txt = ""
                if is_multi_year_subtitle(txt):
                    txt = ""
                if txt:
                    exp_row[c] = txt
                    if c > 0:
                        last_val = txt
                else:
                    if last_val and c > 0 and c in keep_cols:
                        exp_row[c] = last_val
            else:
                if last_val and c > 0 and c in keep_cols:
                    exp_row[c] = last_val
        expanded_header_rows.append(exp_row)

    # Build hierarchical column names
    col_names = []
    for c in keep_cols:
        parts = []
        for r_idx in range(len(expanded_header_rows)):
            val = expanded_header_rows[r_idx][c]
            if val:
                txt = clean_english_cell_text(val)
                if txt and not txt.startswith("Table ") and txt not in parts:
                    parts.append(txt)
        col_header = " - ".join(parts).strip()
        if not col_header:
            col_header = f"Column_{openpyxl.utils.get_column_letter(min_col + c)}"
        col_names.append(col_header)

    seen = {}
    unique_col_names = []
    for name in col_names:
        if name in seen:
            seen[name] += 1
            unique_col_names.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 1
            unique_col_names.append(name)

    # Build raw output rows
    output_rows = []
    for r_idx, r in enumerate(data_rows):
        row_cells = []
        is_empty = True
        
        non_empty_cells = [(c, clean_val(r[c])) for c in keep_cols if r[c] is not None and clean_val(r[c]) != ""]
        if len(non_empty_cells) == 1 and not isinstance(r[non_empty_cells[0][0]], (int, float)):
            txt = clean_english_cell_text(non_empty_cells[0][1])
            if txt and not re.match(r'^-?[\d,.]+$', txt):
                row_cells = [txt if i == 0 else "" for i in range(len(keep_cols))]
                output_rows.append(row_cells)
                continue

        for c in keep_cols:
            val = r[c]
            fname = data_font_rows[r_idx][c]
            if val is not None and not is_pure_dhivehi_cell(val, fname):
                v_clean = clean_english_cell_text(val)
                row_cells.append(v_clean)
                if v_clean:
                    is_empty = False
            else:
                row_cells.append("")
                
        if not is_empty:
            output_rows.append(row_cells)

    # Drop columns where ALL data rows are completely empty
    final_cols_idx = []
    for col_i in range(len(unique_col_names)):
        has_any_data = any(row[col_i] != "" for row in output_rows)
        if has_any_data:
            final_cols_idx.append(col_i)

    final_col_names = [unique_col_names[i] for i in final_cols_idx]
    final_output_rows = [[row[i] for i in final_cols_idx] for row in output_rows]

    meta = {
        "table_id": base_name_clean,
        "title_en": title_en,
        "title_dv": title_dv,
        "unit": unit_info,
        "source": source,
        "notes": notes,
        "columns": final_col_names,
        "row_count": len(final_output_rows)
    }

    os.makedirs(output_dir, exist_ok=True)
    csv_path = os.path.join(output_dir, f"{base_name_clean}.csv")
    json_path = os.path.join(output_dir, f"{base_name_clean}.json")

    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(final_col_names)
        writer.writerows(final_output_rows)

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)

    return meta, csv_path, json_path

def main():
    root_dir = 'raw_stats_data'
    base_output_dir = 'datasets'
    os.makedirs(base_output_dir, exist_ok=True)
    
    chapter_dirs = sorted([d for d in glob.glob(os.path.join(root_dir, '*')) if os.path.isdir(d)])
    
    total_processed = 0
    total_success = 0
    total_failed = 0
    
    start_time = time.time()
    summary_report = []
    
    for ch_path in chapter_dirs:
        ch_name = os.path.basename(ch_path.rstrip(' /')).strip()
        target_out = os.path.join(base_output_dir, ch_name)
        xlsx_files = sorted(glob.glob(os.path.join(ch_path, '*.xlsx')))
        
        print(f"\n==========================================")
        print(f"Processing Chapter: {ch_name} ({len(xlsx_files)} tables)")
        print(f"==========================================")
        
        ch_success = 0
        ch_failed = 0
        
        for f in xlsx_files:
            total_processed += 1
            bname = os.path.basename(f)
            try:
                meta, cp, jp = process_single_table(f, target_out)
                ch_success += 1
                total_success += 1
                print(f"  [OK] {bname:<15} -> {meta['row_count']:>4} rows, {len(meta['columns']):>2} cols | {meta['title_en'][:60]}")
            except Exception as e:
                ch_failed += 1
                total_failed += 1
                print(f"  [ERROR] {bname}: {e}")
                
        summary_report.append({
            "chapter": ch_name,
            "total_files": len(xlsx_files),
            "success": ch_success,
            "failed": ch_failed
        })
        
    elapsed = time.time() - start_time
    print("\n" + "="*50)
    print("FINAL SUMMARY REPORT")
    print("="*50)
    print(f"Total files processed: {total_processed}")
    print(f"Successfully converted: {total_success}")
    print(f"Failed: {total_failed}")
    print(f"Time elapsed: {elapsed:.2f} seconds")
    print("="*50)
    for rep in summary_report:
        print(f"{rep['chapter']:<40} : {rep['success']}/{rep['total_files']} OK")

if __name__ == '__main__':
    main()

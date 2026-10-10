import os
import re
import json
import csv
import openpyxl
from openpyxl.utils import get_column_letter

THAANA_REGEX = re.compile(r'[\u0780-\u07BF]')

LEGACY_FONTS = {'faseyha', 'faruma', 'randhoo', 'kuredhi', 'waheed', 'voguepsmt', 'a_faseyha', 'a_faruma', 'a_randhoo'}

def is_dhivehi_font(font_name):
    if not font_name:
        return False
    fn = font_name.lower()
    return any(lf in fn for lf in LEGACY_FONTS)

def is_pure_dhivehi_cell(val, font_name=''):
    if val is None:
        return False
    val_str = str(val).strip()
    if not val_str:
        return False
    # If it's a number, it's not Dhivehi text
    if isinstance(val, (int, float)) or re.match(r'^-?[\d,.]+$', val_str):
        return False
    if THAANA_REGEX.search(val_str):
        # Check if there is also English text
        has_latin = bool(re.search(r'[A-Za-z]{2,}', val_str))
        if not has_latin:
            return True
    if is_dhivehi_font(font_name):
        return True
    # Common legacy ASCII Dhivehi phrases
    known_legacy = {'clwvwt', 'ulwvwt', 'cnunemihob', 'cscnijed', 'cnehirif', 'cnehcnwa', 'wlcmuj', 'cnwt', 'elwm', 'cnwg .s'}
    if val_str.lower() in known_legacy:
        return True
    return False

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        val_rounded = round(val, 8)
        val_str = f"{val_rounded:f}".rstrip('0').rstrip('.')
        return val_str
    if isinstance(val, int):
        return str(val)
    val_str = str(val).strip()
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

def clean_english_cell_text(val):
    if not val:
        return ""
    val_str = clean_val(val)
    # Remove Thaana characters from mixed cells
    val_str = THAANA_REGEX.sub('', val_str).strip()
    # Remove common legacy dhivehi words attached to english
    val_str = re.sub(r'\b(cnunemihOb|cscnijed|cnehirif|cnehcnwa|wlcmuj|clwvWt|ulwvWt|cnwt)\b', '', val_str, flags=re.IGNORECASE).strip()
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

def extract_table(xlsx_path, output_dir):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    # Identify target sheet
    base_name = os.path.splitext(os.path.basename(xlsx_path))[0].strip()
    base_name_clean = base_name.replace(' ', '')
    
    target_sheet = None
    for name in wb.sheetnames:
        sheet = wb[name]
        if sheet.sheet_state != 'hidden':
            if name.strip().replace(' ', '') == base_name_clean:
                target_sheet = sheet
                break
    if not target_sheet:
        for name in wb.sheetnames:
            sheet = wb[name]
            if sheet.sheet_state != 'hidden':
                target_sheet = sheet
                break
                
    sheet = target_sheet
    
    # Print area bounding if available
    min_col, min_row = 1, 1
    max_col, max_row = sheet.max_column, sheet.max_row
    
    if sheet.print_area:
        pa = sheet.print_area
        if isinstance(pa, list) and pa:
            pa = pa[0]
        m = re.search(r'\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)', str(pa))
        if m:
            c1, r1, c2, r2 = m.groups()
            min_col = openpyxl.utils.column_index_from_string(c1)
            min_row = int(r1)
            max_col = openpyxl.utils.column_index_from_string(c2)
            max_row = int(r2)

    # 1. Expand merged cells so values populate their full visual span
    # We build a 2D grid
    grid = [[None for _ in range(max_col - min_col + 1)] for _ in range(max_row - min_row + 1)]
    font_grid = [["" for _ in range(max_col - min_col + 1)] for _ in range(max_row - min_row + 1)]
    
    for r_idx, r in enumerate(range(min_row, max_row + 1)):
        for c_idx, c in enumerate(range(min_col, max_col + 1)):
            cell = sheet.cell(r, c)
            grid[r_idx][c_idx] = cell.value
            font_grid[r_idx][c_idx] = cell.font.name if cell.font and cell.font.name else ""

    # Propagate merged cell top-left value across merge ranges
    for rng in sheet.merged_cells.ranges:
        r_start = max(min_row, rng.min_row) - min_row
        r_end = min(max_row, rng.max_row) - min_row
        c_start = max(min_col, rng.min_col) - min_col
        c_end = min(max_col, rng.max_col) - min_col
        
        top_val = sheet.cell(rng.min_row, rng.min_col).value
        top_font = sheet.cell(rng.min_row, rng.min_col).font
        fname = top_font.name if top_font and top_font.name else ""
        
        for r_i in range(r_start, r_end + 1):
            for c_i in range(c_start, c_end + 1):
                if 0 <= r_i < len(grid) and 0 <= c_i < len(grid[0]):
                    grid[r_i][c_i] = top_val
                    font_grid[r_i][c_i] = fname

    # 2. Extract Title, Metadata, Source, Notes
    title_en = ""
    title_dv = ""
    unit_info = ""
    notes = []
    source = ""
    
    data_end_row = len(grid)
    header_start_row = -1
    
    for r_idx, row in enumerate(grid):
        non_empty = [(c, clean_val(v), font_grid[r_idx][c]) for c, v in enumerate(row) if v is not None and clean_val(v) != ""]
        if not non_empty:
            continue
        first_c, first_v, first_f = non_empty[0]
        
        # Check title patterns
        if re.search(r'Table\s+\d+(\.\d+)?[a-z]?\s*:', first_v, re.IGNORECASE):
            title_en = first_v
            continue
        if re.search(r'ތާވަލ[ުް]?\s+\d+(\.\d+)?[a-z]?\s*:', first_v) or first_v.endswith('clwvWt') or 'clwvWt :' in first_v or 'ulwvWt :' in first_v:
            title_dv = first_v
            continue
        if re.match(r'^\([A-Za-z0-9\s,%_/\-]+\)$', first_v) and not title_en:
            unit_info = first_v
            continue
            
        # Notes & Sources
        lower_v = first_v.lower()
        if any(kw in lower_v for kw in ['source:', 'source :', 'source ;', 'ﻤﻌﻠﻮﻣﺎﺕ', 'މައުލޫމާތު']):
            source = first_v
            if data_end_row == len(grid):
                data_end_row = r_idx
            continue
        if any(lower_v.startswith(kw) for kw in ['note:', 'note :', '* note', '1_/', '2_/', '3_/', '4_/', '5_/']):
            notes.append(first_v)
            if data_end_row == len(grid):
                data_end_row = r_idx
            continue
        if data_end_row < len(grid):
            # Already in notes section
            notes.append(first_v)

    if not title_en:
        for r_idx in range(min(5, len(grid))):
            for c_idx, v in enumerate(grid[r_idx]):
                if v and isinstance(v, str) and 'Table' in v:
                    title_en = clean_val(v)
                    break
            if title_en:
                break

    # Determine Header rows vs Data rows
    for r_idx in range(min(12, data_end_row)):
        cells = [clean_val(v) for v in grid[r_idx] if v is not None and clean_val(v) != ""]
        if not cells:
            continue
        line_str = " ".join(cells)
        if title_en and title_en in line_str:
            continue
        if title_dv and title_dv in line_str:
            continue
        if header_start_row == -1:
            header_start_row = r_idx
            break

    # Find where data rows start: first row with numbers or repeated category rows
    data_start_row = header_start_row + 1
    for r_idx in range(header_start_row, data_end_row):
        row = grid[r_idx]
        numeric_count = sum(1 for v in row if isinstance(v, (int, float)) or (isinstance(v, str) and re.match(r'^-?[\d,.]+$', v.strip())))
        if numeric_count >= 1:
            data_start_row = r_idx
            break

    header_rows = grid[header_start_row:data_start_row]
    header_fonts = font_grid[header_start_row:data_start_row]
    raw_data_rows = grid[data_start_row:data_end_row]
    raw_data_fonts = font_grid[data_start_row:data_end_row]

    # Filter out Pure Dhivehi data rows (e.g. Dhivehi month/age names placed on alternating rows)
    data_rows = []
    data_font_rows = []
    for r_idx, r in enumerate(raw_data_rows):
        has_num = any(isinstance(v, (int, float)) or (isinstance(v, str) and re.match(r'^-?[\d,.]+$', v.strip())) for v in r if v is not None)
        non_empty = [v for v in r if v is not None and clean_val(v) != ""]
        if not non_empty:
            continue
        if not has_num:
            # Check if all cells in this row are Dhivehi
            all_dv = all(is_pure_dhivehi_cell(v, raw_data_fonts[r_idx][c]) for c, v in enumerate(r) if v is not None and clean_val(v) != "")
            if all_dv:
                continue
        data_rows.append(r)
        data_font_rows.append(raw_data_fonts[r_idx])

    # Forward-fill English header labels horizontally across empty/Dhivehi columns
    # for each header row (e.g. 1995 Census spanning B, C, D)
    num_cols = len(grid[0]) if grid else 0
    expanded_header_rows = []
    for h_idx, h_row in enumerate(header_rows):
        exp_row = list(h_row)
        last_en_val = None
        for c in range(num_cols):
            v = h_row[c]
            f = header_fonts[h_idx][c]
            if v is not None and not is_pure_dhivehi_cell(v, f) and clean_english_cell_text(v):
                # Don't use title as column header
                t_clean = clean_english_cell_text(v)
                if title_en and t_clean.lower() == title_en.lower():
                    exp_row[c] = None
                    last_en_val = None
                else:
                    last_en_val = t_clean
                    exp_row[c] = last_en_val
            else:
                if is_pure_dhivehi_cell(v, f):
                    exp_row[c] = None
                elif last_en_val is not None:
                    # check if this column has data or subheader below
                    has_col_content = any(dr[c] is not None for dr in data_rows) or any(hr[c] is not None for hr in header_rows)
                    if has_col_content:
                        exp_row[c] = last_en_val
        expanded_header_rows.append(exp_row)

    # Find columns to keep:
    # Must have non-empty data or header, and NOT be a pure Dhivehi column
    keep_cols = []
    for c in range(num_cols):
        col_data = [r[c] for r in data_rows if r[c] is not None and clean_val(r[c]) != ""]
        col_header = [clean_val(r[c]) for r in header_rows if r[c] is not None and clean_val(r[c]) != ""]
        if not col_data and not col_header:
            continue
            
        # Check if all cells in this column are pure Dhivehi text
        all_col_cells = [(r[c], data_font_rows[i][c]) for i, r in enumerate(data_rows) if r[c] is not None and clean_val(r[c]) != ""]
        all_col_cells += [(r[c], header_fonts[i][c]) for i, r in enumerate(header_rows) if r[c] is not None and clean_val(r[c]) != ""]
        
        is_col_dhivehi = True
        for val, fname in all_col_cells:
            v_str = clean_val(val)
            if isinstance(val, (int, float)) or re.match(r'^-?[\d,.]+$', v_str):
                is_col_dhivehi = False
                break
            if not is_pure_dhivehi_cell(val, fname):
                is_col_dhivehi = False
                break
                
        if not is_col_dhivehi:
            keep_cols.append(c)

    # Check for duplicate consecutive columns (like Col 2 Locality and Col 3 Locality in 1.5)
    filtered_keep_cols = []
    for i, c in enumerate(keep_cols):
        if i > 0:
            prev_c = keep_cols[i - 1]
            # Compare data values
            prev_vals = [clean_val(r[prev_c]) for r in data_rows]
            curr_vals = [clean_val(r[c]) for r in data_rows]
            if prev_vals and curr_vals and prev_vals == curr_vals:
                # Exact duplicate column, skip it
                continue
        filtered_keep_cols.append(c)
    keep_cols = filtered_keep_cols

    # Build hierarchical English column names from expanded_header_rows
    col_names = []
    for c in keep_cols:
        parts = []
        for r_idx in range(len(expanded_header_rows)):
            val = expanded_header_rows[r_idx][c]
            if val:
                v_clean = clean_english_cell_text(val)
                # Strip table title if contained
                if title_en and v_clean.startswith("Table "):
                    continue
                if v_clean and v_clean not in parts:
                    parts.append(v_clean)
        col_header = " - ".join(parts).strip()
        if not col_header:
            col_header = f"Column_{openpyxl.utils.get_column_letter(min_col + c)}"
        col_names.append(col_header)

    # Clean duplicates in col_names if any
    seen = {}
    unique_col_names = []
    for name in col_names:
        if name in seen:
            seen[name] += 1
            unique_col_names.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 1
            unique_col_names.append(name)

    # Build clean output rows
    output_rows = []
    for r_idx, r in enumerate(data_rows):
        row_cells = []
        is_empty = True
        
        # Section header check (e.g. 'North Thiladhunmathi (HA)')
        non_empty_cells = [(c, clean_val(r[c])) for c in keep_cols if r[c] is not None and clean_val(r[c]) != ""]
        if len(non_empty_cells) == 1 and not isinstance(r[non_empty_cells[0][0]], (int, float)):
            txt = clean_english_cell_text(non_empty_cells[0][1])
            if txt and not re.match(r'^-?[\d,.]+$', txt):
                row_cells = [txt if i == 0 else "" for i in range(len(keep_cols))]
                output_rows.append(row_cells)
                continue

        for c in keep_cols:
            val = r[c]
            fname = data_font_rows[r_idx][c]
            if val is not None and not is_pure_dhivehi_cell(val, fname):
                v_clean = clean_english_cell_text(val)
                row_cells.append(v_clean)
                if v_clean:
                    is_empty = False
            else:
                row_cells.append("")
                
        if not is_empty:
            output_rows.append(row_cells)

    # Prepare metadata dict
    meta = {
        "table_id": base_name_clean,
        "title_en": title_en,
        "title_dv": title_dv,
        "unit": unit_info,
        "source": source,
        "notes": notes,
        "columns": unique_col_names,
        "row_count": len(output_rows)
    }

    os.makedirs(output_dir, exist_ok=True)
    csv_path = os.path.join(output_dir, f"{base_name_clean}.csv")
    json_path = os.path.join(output_dir, f"{base_name_clean}.json")

    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(unique_col_names)
        writer.writerows(output_rows)

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)

    return meta, csv_path, json_path

if __name__ == '__main__':
    test_files = [
        'raw_stats_data/1.0 Geography /1.1.xlsx',
        'raw_stats_data/1.0 Geography /1.2.xlsx',
        'raw_stats_data/1.0 Geography /1.3.xlsx',
        'raw_stats_data/1.0 Geography /1.4.xlsx',
        'raw_stats_data/1.0 Geography /1.5.xlsx',
        'raw_stats_data/3.0 population /3.1.xlsx',
        'raw_stats_data/3.0 population /3.2.xlsx',
        'raw_stats_data/3.0 population /3.3.xlsx',
        'raw_stats_data/3.0 population /3.4.xlsx',
        'raw_stats_data/3.0 population /3.5.xlsx'
    ]
    
    out_dir = 'datasets/sample_test'
    for f in test_files:
        try:
            m, c_p, j_p = extract_table(f, out_dir)
            print(f"[OK] {f} -> {c_p} ({m['row_count']} rows, {len(m['columns'])} cols)")
            print(f"     Title: {m['title_en']}")
            print(f"     Cols: {m['columns'][:5]}")
        except Exception as e:
            import traceback
            print(f"[FAIL] {f}: {e}")
            traceback.print_exc()
